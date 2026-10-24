# -*- coding: utf-8 -*-
"""
실무용 템플릿 채우기
사찰 넷제로 진단 도구 - 불교탄소중립실천단 지수광풍

「온실가스 배출량 산정 실무용 템플릿」의 노란 셀(활동량)만 채웁니다.
수식·참조표·서식은 손대지 않으므로, 엑셀에서 열면 5단계 계산이 그대로 돕니다.

  ③Scope1_고정연소   B8:B15 연료명 / D8:D15 활동량
  ④Scope1_이동연소   B8:B14 연료명 / D8:D14 활동량
  ⑤Scope2_전력       D7 구매전력(kWh) / D10 재생에너지(kWh)
"""
from __future__ import annotations

import datetime as dt
import os
import shutil
from dataclasses import dataclass
from typing import List, Optional, Sequence, Tuple

from openpyxl import load_workbook

import emissions as em

# 지수광풍판 — 작업지시 06 으로 개정한 서식입니다.
# 구판은 `..._구판.xlsx` 로 남겨 두었습니다(옛 결과를 다시 볼 일이 있으면).
TEMPLATE_NAME = "서식_온실가스_배출량_산정_템플릿_지수광풍판.xlsx"
TEMPLATE_BLANK = "서식_온실가스_배출량_산정_템플릿_지수광풍판_빈서식.xlsx"

SH_FIXED = "③Scope1_고정연소"
SH_MOBILE = "④Scope1_이동연소"
SH_POWER = "⑤Scope2_전력"
SH_SUM = "⑥종합요약"
SH_INFO = "⓪기본정보"
SH_SRC = "⑧조회근거"
SRC_ROWS = range(14, 20)      # ⑧ 의 지번별 줄 (6칸)
BIO_COL = 20                  # ③ 의 생물기원 CO2 열 (T). 합계에 안 들어갑니다
SOLAR_ROWS = (21, 22)         # ⑤ 의 자가 태양광 발전량 / 자가소비량

FIXED_ROWS = range(8, 16)     # 8칸
MOBILE_ROWS = range(8, 15)    # 7칸


class TemplateError(Exception):
    """화면에 그대로 보여줄 수 있는 오류."""


@dataclass
class Activity:
    """템플릿에 넣을 활동량 한 줄.

    같은 연료라도 보일러에 쓰면 고정연소, 차량에 쓰면 이동연소입니다.
    **계수도 다릅니다**(2026-10-24 개정). 고정연소는 별표10 상업·공공 열에
    별표6 산화계수, 도로는 별표6 <표-4> 에 산화계수 항이 없습니다.
    들어갈 시트도 다릅니다(③ / ④).
    scope 를 비워 두면 흔한 쓰임으로 정합니다.
    """

    fuel: str
    qty: float
    scope: str = ""
    note: str = ""

    def __post_init__(self):
        if not self.scope:
            src = em.FUELS.get(self.fuel) or em.ROAD.get(self.fuel)
            if src is None:
                raise TemplateError(f"참조표에 없는 연료입니다: {self.fuel}")
            self.scope = src.default_scope
        if self.scope not in ("고정연소", "이동연소"):
            raise TemplateError(f"고정연소 또는 이동연소여야 합니다: {self.scope}")
        if self.fuel not in em.table_for(self.scope):
            where = "도로 이동연소" if self.scope == "이동연소" else "고정연소"
            raise TemplateError(f"{where} 참조표에 없는 연료입니다: {self.fuel}")

    @property
    def unit(self) -> str:
        return em.table_for(self.scope)[self.fuel].unit


def find_template(folder: str = "") -> str:
    """서식 파일을 찾습니다."""
    here = folder or os.path.dirname(os.path.abspath(__file__))
    path = os.path.join(here, TEMPLATE_NAME)
    if not os.path.exists(path):
        raise TemplateError(
            f"서식 파일이 없습니다: {TEMPLATE_NAME}\n"
            "진단 도구와 같은 폴더에 두어야 합니다."
        )
    return path


def _write_rows(ws, rows: Sequence[int], items: List[Activity], label: str) -> None:
    if len(items) > len(rows):
        raise TemplateError(
            f"{label} 칸이 {len(rows)}개인데 {len(items)}개를 넣으려 합니다. "
            "템플릿에서 행을 늘린 뒤 다시 해주세요."
        )
    bio = label == "고정연소"
    for i, r in enumerate(rows):
        if i < len(items):
            a = items[i]
            ws.cell(row=r, column=2).value = a.fuel           # B 연료명
            ws.cell(row=r, column=4).value = a.qty            # D 활동량
            if bio:
                # 생물기원 CO2 는 합계에 넣지 않고 따로 적습니다(고시 제18조).
                b = em.burn(a.fuel, a.qty, a.scope)
                ws.cell(row=r, column=BIO_COL).value = (
                    round(b.bio_co2_kg, 1) if b.bio_co2_kg else None)
        else:
            ws.cell(row=r, column=2).value = None
            ws.cell(row=r, column=4).value = None
            if bio:
                ws.cell(row=r, column=BIO_COL).value = None


def fill(
    out_path: str,
    temple: str,
    year: int,
    elec_kwh: float,
    activities: Sequence[Activity] = (),
    renewable_kwh: float = 0.0,
    template_path: Optional[str] = None,
    note: str = "",
    road_addr: str = "",
    jibun_addr: str = "",
    solar_kwh: float = 0.0,
    solar_self_kwh: float = 0.0,
    sources: Sequence[dict] = (),
    merge_basis: str = "",
) -> str:
    """sources 는 지번마다 한 줄입니다(⑧조회근거).

        {"지번": "엄사리 375-3", "연면적": 336.9, "주차장제외": 336.9,
         "전기": 25699, "가스": 0, "동수": 1, "전통": 0, "비고": ""}

    merge_basis 는 묶음 판정입니다 — 공개자료_전지번 / 고지서확인 / 모름.

    **「비고」는 회원이 적을 때만 채웁니다. 기본은 빈칸입니다.**
    대장 주용도를 보고 「법당」·「요사채」 같은 이름을 **붙이지 마십시오.**
    원오사에서 법당이 단독주택, 숙소가 제2종근린생활시설로 올라가 있었습니다
    (CLAUDE.md 7장 「대장 주용도로 건물의 쓰임을 짐작하지 말 것」).
    """
    """템플릿 사본을 만들어 활동량을 채워 넣습니다. 만들어진 파일 경로를 돌려줍니다."""
    src = template_path or find_template()
    shutil.copyfile(src, out_path)

    wb = load_workbook(out_path)          # 수식 보존 (data_only=False 기본값)
    for need in (SH_FIXED, SH_MOBILE, SH_POWER):
        if need not in wb.sheetnames:
            raise TemplateError(f"서식에 '{need}' 시트가 없습니다. 서식 파일을 확인해 주세요.")

    fixed = [a for a in activities if a.scope == "고정연소"]
    mobile = [a for a in activities if a.scope == "이동연소"]

    _write_rows(wb[SH_FIXED], FIXED_ROWS, fixed, "고정연소")
    _write_rows(wb[SH_MOBILE], MOBILE_ROWS, mobile, "이동연소")

    ws = wb[SH_POWER]
    ws["D7"] = float(elec_kwh or 0)
    ws["D10"] = float(renewable_kwh or 0)
    ws.cell(SOLAR_ROWS[0], 4).value = float(solar_kwh or 0) or None
    ws.cell(SOLAR_ROWS[1], 4).value = float(solar_self_kwh or 0) or None

    stamp = dt.date.today().strftime("%Y. %m. %d.")
    if SH_SRC in wb.sheetnames:
        _fill_sources(wb[SH_SRC], temple, year, road_addr, stamp,
                      sources or [{"지번": jibun_addr}], merge_basis, activities)

    if SH_INFO in wb.sheetnames:
        info = wb[SH_INFO]
        info["C5"] = temple
        info["C6"] = road_addr
        info["C7"] = jibun_addr
        info["C8"] = f"{year}년 사용량"
        info["C9"] = stamp
    tail = f"  |  {note}" if note else ""
    wb[SH_SUM]["B4"] = (
        f"대상: {temple}  |  기준연도: {year}년  |  작성일: {stamp}{tail}"
    )
    wb[SH_SUM]["B5"] = (
        "활동량은 사찰 넷제로 진단 도구가 채웠습니다. "
        "전기·도시가스는 건축HUB 건물에너지 공개자료, 나머지는 현장 조사값입니다."
    )

    wb.save(out_path)
    return out_path


def from_diagnosis(
    out_path: str,
    temple: str,
    year: int,
    elec_kwh: float,
    gas_kwh: float,
    gas_basis: str = em.GAS_BASIS_DEFAULT,
    extra: Sequence[Activity] = (),
    **kw,
) -> Tuple[str, List[Activity]]:
    """개략 진단 결과를 그대로 템플릿에 넣습니다.

    건물에너지 API 의 가스는 kWh 라서 Nm3 으로 바꿔 넣습니다.
    extra 에 현장에서 확인한 LPG·등유·차량연료를 넘기면 함께 들어갑니다.
    """
    acts: List[Activity] = []
    if gas_kwh:
        nm3 = em.gas_kwh_to_nm3(gas_kwh, gas_basis)
        acts.append(Activity("도시가스/LNG", round(nm3, 1), scope="고정연소",
                             note=f"{gas_kwh:,.0f} kWh 를 {gas_basis} 기준으로 환산"))
    acts.extend(extra)

    path = fill(out_path, temple, year, elec_kwh, acts,
                note=f"도시가스 환산기준 {gas_basis}", **kw)
    return path, acts


def preview(elec_kwh: float, activities: Sequence[Activity],
            elec_basis: str = em.ELEC_DEFAULT, gwp: str = em.GWP_DEFAULT) -> dict:
    """템플릿이 낼 값을 미리 계산해 화면에 보여줍니다. (엑셀 없이 확인용)"""
    rows = []
    s1 = 0.0
    bio = 0.0          # 생물기원 CO2 — 합계에 넣지 않고 따로 보여 줍니다
    for a in activities:
        b = em.burn(a.fuel, a.qty, a.scope)
        t = b.ar5_t if gwp == "AR5" else b.ar6_t
        s1 += t
        bio += b.bio_co2_kg
        rows.append({"연료": a.fuel, "구분": a.scope, "활동량": f"{a.qty:,.1f} {a.unit}",
                     "발열량": f"{b.hv} MJ/{b.unit}", "TJ": f"{b.tj:.6f}",
                     "배출량": f"{t:.3f} t", "산식": b.formula(gwp)})
    s2 = em.power_t(elec_kwh, elec_basis)
    return {"rows": rows, "scope1_t": s1, "scope2_t": s2, "total_t": s1 + s2,
            "bio_co2_kg": bio, "elec_basis": elec_basis, "gwp": gwp}


def _fill_sources(ws, temple, year, road_addr, stamp, sources,
                  merge_basis, activities) -> None:
    """⑧조회근거 를 채웁니다. 숫자가 어디서 왔는지 남기는 자리입니다."""
    ws["C5"] = temple
    ws["C6"] = road_addr
    ws["C7"] = f"{year}년 사용량"
    ws["C8"] = f"{stamp} 조회"
    ws["C10"] = {
        "공개자료_전지번": "모든 지번에 전기 기록이 있어 합산했습니다 (고지서 확인 불필요)",
        "고지서확인": "고지서로 한전 계약이 하나임을 확인했습니다",
        "모름": "확인되지 않았습니다 — 전기 0 인 지번이 끼어 있습니다. "
                "고지서의 고객번호 개수와 사용 장소 주소를 확인해 주세요",
    }.get(merge_basis, "지번 하나만 조회했습니다 (묶음 없음)")

    if len(sources) > len(SRC_ROWS):
        raise TemplateError(
            f"⑧조회근거 칸이 {len(SRC_ROWS)}개인데 지번 {len(sources)}개를 "
            "넣으려 합니다.")
    for i, r in enumerate(SRC_ROWS):
        src = sources[i] if i < len(sources) else None
        for j, key in enumerate(("지번", "연면적", "주차장제외", "전기",
                                 "가스", "동수", "전통", "비고")):
            # 「비고」는 받은 것만 적습니다. 도구가 만들어 내지 않습니다.
            ws.cell(r, 2 + j).value = (src or {}).get(key) if src else None

    # 현장 입력값 — 안 넣었으면 「미입력」이라고 적습니다.
    typed = [f"{a.fuel} {a.qty:,.0f} {a.unit}" for a in activities
             if a.qty and a.fuel != "도시가스/LNG"]
    ws["B23"] = " · ".join(typed) if typed else (
        "미입력 — LPG·등유·화목·차량 연료를 넣지 않았습니다. "
        "아래 「반영되지 않은 것」을 함께 보십시오.")
